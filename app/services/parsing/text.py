"""文本抽取：按文件类型产出带位置标识的 TextBlock 列表（设计文档 §7.4.2–§7.4.3）。"""

from __future__ import annotations

import logging
import time
from dataclasses import dataclass
from pathlib import Path

from app.config import Settings
from app.services.parsing.ocr import OcrUnavailableError, ocr_pil_image

logger = logging.getLogger(__name__)

_PDF_MAGIC = b"%PDF"
_PNG_MAGIC = b"\x89PNG\r\n\x1a\n"
_JPG_MAGIC = b"\xff\xd8\xff"
_ZIP_MAGIC = b"PK\x03\x04"

# 页面文本少于该字符数时尝试 OCR
_OCR_TEXT_THRESHOLD = 20
# 200 DPI 渲染（PDF 默认 72 DPI）
_OCR_RENDER_SCALE = 200 / 72


class ParseFailure(Exception):
    """解析失败（带 blocked_reason 语义的错误码：DOCUMENT_EMPTY / PARSE_FAILED / OCR_FAILED）。"""

    def __init__(self, code: str, message: str) -> None:
        self.code = code
        self.message = message
        super().__init__(f"{code}: {message}")


@dataclass(frozen=True)
class TextBlock:
    page_no: int | None
    position: str
    text: str
    source: str  # native | ocr
    confidence: float = 1.0


@dataclass
class DocumentText:
    blocks: list[TextBlock]
    page_count: int | None
    ocr_used: bool


def detect_file_type(path: Path, declared: str | None = None) -> str:
    """魔数优先、扩展名兜底。"""
    head = path.read_bytes()[:8]
    if head.startswith(_PDF_MAGIC):
        return "pdf"
    if head.startswith(_PNG_MAGIC):
        return "png"
    if head.startswith(_JPG_MAGIC):
        return "jpg"
    if head.startswith(_ZIP_MAGIC):
        return "docx"  # docx 是 zip；进一步的 [Content_Types].xml 校验由 python-docx 完成
    if declared:
        return declared.lower()
    raise ParseFailure("PARSE_FAILED", f"无法识别的文件类型：{path.name}")


def _ensure_deadline(deadline: float) -> None:
    """超过解析预算（PARSE_TIMEOUT_SECONDS）时中止并给出可诊断错误。"""
    if time.monotonic() > deadline:
        raise ParseFailure("PARSE_FAILED", "解析超时（超过 PARSE_TIMEOUT_SECONDS）")


def extract_document(path: Path, file_type: str, settings: Settings) -> DocumentText:
    detected = detect_file_type(path, file_type)
    deadline = time.monotonic() + max(1.0, float(settings.parse_timeout_seconds))
    if detected == "pdf":
        return _extract_pdf(path, settings, deadline)
    if detected == "docx":
        return _extract_docx(path, deadline)
    if detected in {"png", "jpg", "jpeg"}:
        return _extract_image(path, settings, deadline)
    raise ParseFailure("PARSE_FAILED", f"不支持的文件类型：{detected}")


def _extract_pdf(path: Path, settings: Settings, deadline: float) -> DocumentText:
    import pypdfium2 as pdfium

    _ensure_deadline(deadline)
    pdf = pdfium.PdfDocument(str(path))
    try:
        page_count = len(pdf)
        if page_count > settings.max_pages:
            raise ParseFailure("PARSE_FAILED", f"页数超限：{page_count} > {settings.max_pages}")

        blocks: list[TextBlock] = []
        ocr_used = False
        for page_index in range(page_count):
            _ensure_deadline(deadline)
            page = pdf[page_index]
            text = _pdf_page_text(page)
            if len(text) >= _OCR_TEXT_THRESHOLD:
                for para_index, para in enumerate(_split_paragraphs(text)):
                    blocks.append(
                        TextBlock(
                            page_no=page_index + 1,
                            position=f"第{page_index + 1}页/第{para_index + 1}段",
                            text=para,
                            source="native",
                        )
                    )
                continue

            image = page.render(scale=_OCR_RENDER_SCALE).to_pil()
            ocr_text, ocr_confidence = ocr_pil_image(image)
            _ensure_deadline(deadline)
            ocr_text = ocr_text.strip()
            if ocr_text:
                ocr_used = True
                for line_index, line in enumerate(_split_paragraphs(ocr_text)):
                    blocks.append(
                        TextBlock(
                            page_no=page_index + 1,
                            position=f"第{page_index + 1}页/OCR第{line_index + 1}行",
                            text=line,
                            source="ocr",
                            confidence=ocr_confidence,
                        )
                    )
            elif text.strip():
                blocks.append(
                    TextBlock(page_index + 1, f"第{page_index + 1}页", text.strip(), "native")
                )
        return DocumentText(blocks=blocks, page_count=page_count, ocr_used=ocr_used)
    finally:
        pdf.close()


def _pdf_page_text(page) -> str:
    textpage = page.get_textpage()
    try:
        return (textpage.get_text_bounded() or "").strip()
    except AttributeError:  # pragma: no cover - 兼容旧版 API
        return (textpage.get_text_range() or "").strip()


def _extract_docx(path: Path, deadline: float) -> DocumentText:
    import docx

    _ensure_deadline(deadline)
    document = docx.Document(str(path))
    blocks: list[TextBlock] = []
    for index, paragraph in enumerate(document.paragraphs):
        if index % 20 == 0:
            _ensure_deadline(deadline)
        text = paragraph.text.strip()
        if text:
            blocks.append(
                TextBlock(page_no=None, position=f"第{index + 1}段", text=text, source="native")
            )
    for table_index, table in enumerate(document.tables):
        for row_index, row in enumerate(table.rows):
            if row_index % 20 == 0:
                _ensure_deadline(deadline)
            cells = [cell.text.strip() for cell in row.cells]
            line = " | ".join(cell for cell in cells if cell)
            if line:
                blocks.append(
                    TextBlock(
                        page_no=None,
                        position=f"表格{table_index + 1}第{row_index + 1}行",
                        text=line,
                        source="native",
                    )
                )
    return DocumentText(blocks=blocks, page_count=None, ocr_used=False)


def _extract_image(path: Path, settings: Settings, deadline: float) -> DocumentText:
    from PIL import Image

    del settings  # 图片路径暂不依赖其它设置，保留签名一致性
    _ensure_deadline(deadline)
    try:
        with Image.open(path) as image:
            image = image.convert("RGB")
    except Exception as exc:  # noqa: BLE001 - 图片损坏统一归为解析失败
        raise ParseFailure("PARSE_FAILED", f"图片无法打开：{exc}") from exc

    ocr_text, ocr_confidence = ocr_pil_image(image)
    _ensure_deadline(deadline)
    ocr_text = ocr_text.strip()
    if not ocr_text:
        raise ParseFailure("OCR_FAILED", "扫描件未识别到任何文字")

    blocks = [
        TextBlock(
            page_no=1,
            position=f"第1页/OCR第{index + 1}行",
            text=line,
            source="ocr",
            confidence=ocr_confidence,
        )
        for index, line in enumerate(_split_paragraphs(ocr_text))
    ]
    return DocumentText(blocks=blocks, page_count=1, ocr_used=True)


def _split_paragraphs(text: str) -> list[str]:
    return [line.strip() for line in text.split("\n") if line.strip()]


__all__ = [
    "DocumentText",
    "OcrUnavailableError",
    "ParseFailure",
    "TextBlock",
    "detect_file_type",
    "extract_document",
]
